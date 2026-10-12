# ============================================================
# generador_docx.py — Plantillas de horario (python-docx).
# Membrete institucional, tabla por días con estilos de la
# plataforma y marco de página verde.
# ============================================================
import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import tema as T

# Paleta (mismos hex que el frontend / PDF)
VERDE = '1A6B35'
VERDE_OSCURO = '0F4521'
VERDE_MUTE = 'E8F5ED'
FILA_ALT = 'F4F9F6'
HORA_BG = 'F0F4F2'
GRIS_BORDE = 'DCE6E0'
GRIS_TEXTO = '4A6357'
GRIS_CLARO = '94A89D'
GRIS_OSCURO = '1E2E27'
ROJO = 'C0392B'
ROJO_BG = 'FEE2E2'
NARANJA = 'F37238'
BLANCO = 'FFFFFF'

FUENTE_TITULO = 'Georgia'      # serif académica (universal en Office)
FUENTE_TEXTO = 'Calibri'       # sans de cuerpo (universal en Office)


def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


def _run(par, texto, *, size=10, bold=False, italic=False,
         color='000000', font=FUENTE_TEXTO):
    run = par.add_run(texto)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)
    return run


def _fondo_celda(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def _borde_celda(cell, hex_color=GRIS_BORDE, sz=4):
    """Aplica borde simple a los 4 lados de la celda."""
    tc_pr = cell._tc.get_or_add_tcPr()
    bordes = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), hex_color)
        bordes.append(el)
    tc_pr.append(bordes)


def _borde_inferior_celda(cell, hex_color=VERDE, sz=12):
    tc_pr = cell._tc.get_or_add_tcPr()
    bordes = OxmlElement('w:tcBorders')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), hex_color)
    bordes.append(el)
    tc_pr.append(bordes)


def _margenes_celda(cell, top=60, bottom=60, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for lado, valor in (('top', top), ('bottom', bottom),
                        ('left', left), ('right', right)):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:w'), str(valor))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tc_pr.append(mar)


def _repetir_encabezado(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def _marco_pagina(section, hex_color=VERDE):
    """Borde de página institucional (fuera del área de texto)."""
    sect_pr = section._sectPr
    pg_bordes = OxmlElement('w:pgBorders')
    pg_bordes.set(qn('w:offsetFrom'), 'page')
    for lado in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:space'), '18')
        el.set(qn('w:color'), hex_color)
        pg_bordes.append(el)
    sect_pr.append(pg_bordes)


def _campo_pagina(par):
    """Agrega el campo dinámico Página (fldSimple PAGE)."""
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    texto = OxmlElement('w:t')
    texto.text = '1'
    run.append(texto)
    fld.append(run)
    par._p.append(fld)


# ------------------------------------------------------------
def calcular_periodos(hora_inicio, hora_fin, duracion_bloque):
    """Paridad con el cálculo del docxGenerator en JS."""
    periodos = []
    h1, m1 = (int(x) for x in hora_inicio.split(':'))
    h2, m2 = (int(x) for x in hora_fin.split(':'))
    actual = h1 * 60 + m1
    limite = h2 * 60 + m2
    while actual < limite:
        inicio = f'{actual // 60:02d}:{actual % 60:02d}'
        actual += int(duracion_bloque)
        fin = f'{actual // 60:02d}:{actual % 60:02d}'
        periodos.append({'inicio': inicio, 'fin': fin, 'receso': False})
    return periodos


def _membrete(doc, institucion):
    """Encabezado con logo + identidad + filete verde (norma membretada)."""
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    # limpiar párrafo inicial vacío
    for par in list(header.paragraphs):
        par._element.getparent().remove(par._element)

    tabla = header.add_table(rows=1, cols=2, width=Cm(17.4))
    tabla.autofit = False
    celda_logo, celda_texto = tabla.rows[0].cells
    celda_logo.width = Cm(2.4)
    celda_texto.width = Cm(15.0)

    par_logo = celda_logo.paragraphs[0]
    par_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        par_logo.add_run().add_picture(T.RUTA_LOGO, width=Cm(2.1))
    except Exception:
        _run(par_logo, 'CECyTE', size=14, bold=True, color=VERDE,
             font=FUENTE_TITULO)

    nombre = institucion.get('nombre', 'CECyTE Plantel 1')
    sede = (f"{institucion.get('sede', 'Poblado Aquiles Serdán, Macuspana')} · "
            f"{institucion.get('estado', 'Tabasco')}")
    lema = institucion.get('lema', 'Orgullo Lince')

    par_nom = celda_texto.paragraphs[0]
    par_nom.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    par_nom.paragraph_format.space_after = Pt(1)
    _run(par_nom, nombre, size=13, bold=True, color=VERDE, font=FUENTE_TITULO)

    par_sede = celda_texto.add_paragraph()
    par_sede.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    par_sede.paragraph_format.space_after = Pt(1)
    _run(par_sede, sede, size=9, color=GRIS_TEXTO)

    par_lema = celda_texto.add_paragraph()
    par_lema.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(par_lema, lema.upper(), size=8, bold=True, color=NARANJA)

    for celda in (celda_logo, celda_texto):
        _borde_inferior_celda(celda, VERDE, sz=12)
        _margenes_celda(celda, top=0, bottom=80, left=60, right=60)

    # separador tras el membrete
    par = header.add_paragraph()
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)


def _pie_pagina(doc, institucion):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(par, (f"{institucion.get('nombre', 'CECyTE Plantel 1')} · "
               f"{institucion.get('sede', 'Macuspana')} · Página "),
         size=8, color=GRIS_CLARO)
    _campo_pagina(par)


def _tira_info(doc, datos):
    """Tira de metadatos (badges) bajo el título."""
    campos = [
        ('SEMESTRE', f"{datos.get('semestre') or '—'}°"),
        ('CICLO', datos.get('ciclo') or '—'),
        ('ESPECIALIDAD', datos.get('especialidad') or '—'),
        ('TURNO', datos.get('turno') or '—'),
    ]
    tabla = doc.add_table(rows=1, cols=len(campos))
    tabla.autofit = False
    ancho_celda = Cm(17.0 / len(campos))
    for celda, (etiqueta, valor) in zip(tabla.rows[0].cells, campos):
        celda.width = ancho_celda
        _fondo_celda(celda, VERDE_MUTE)
        _borde_celda(celda, GRIS_BORDE, sz=4)
        _margenes_celda(celda, top=70, bottom=70, left=140, right=140)
        celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = celda.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        _run(par, etiqueta, size=7.5, bold=True, color=VERDE)
        par2 = celda.add_paragraph()
        par2.paragraph_format.space_after = Pt(0)
        _run(par2, str(valor), size=10, color=GRIS_OSCURO)
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)


def generar_horario(payload):
    """Plantilla DOCX de horario. Devuelve bytes DOCX."""
    inst = payload.get('institucion') or {}
    data = payload.get('data') or payload

    tipo = data.get('tipo', 'grupo')
    nombre = data.get('nombre', '')
    semestre = data.get('semestre', '')
    ciclo = data.get('ciclo', '')
    especialidad = data.get('especialidad', '')
    turno = data.get('turno', 'Matutino')
    dias = data.get('dias') or ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes']
    periodos = data.get('periodos')
    hora_inicio = data.get('hora_inicio', '07:00')
    hora_fin = data.get('hora_fin', '14:10')
    duracion = data.get('duracion_bloque', 50)

    if periodos and isinstance(periodos, list) and periodos:
        horarios = periodos
    else:
        horarios = calcular_periodos(hora_inicio, hora_fin, duracion)
        if not horarios:
            horarios = []
            inicio = 7 * 60
            for _ in range(7):
                horarios.append({
                    'inicio': f'{inicio // 60:02d}:{inicio % 60:02d}',
                    'fin': f'{(inicio + duracion) // 60:02d}:{(inicio + duracion) % 60:02d}',
                    'receso': False,
                })
                inicio += duracion

    titulos = {
        'grupo': f'HORARIO DE GRUPO — {nombre}',
        'maestro': f'HORARIO DE MAESTRO — {nombre}',
        'laboratorio': f'HORARIO DE LABORATORIO — {nombre}',
    }
    titulo = titulos.get(tipo, titulos['grupo'])
    subtitulo = (f"Semestre: {semestre}° · Ciclo: {ciclo} · "
                 f"Especialidad: {especialidad} · Turno: {turno} · "
                 f"Grupo: {nombre}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    _membrete(doc, inst)
    _pie_pagina(doc, inst)
    _marco_pagina(section)

    # Título + subtítulo
    par_titulo = doc.add_paragraph()
    par_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_titulo.paragraph_format.space_before = Pt(6)
    par_titulo.paragraph_format.space_after = Pt(2)
    _run(par_titulo, titulo.upper(), size=15, bold=True, color=VERDE_OSCURO,
         font=FUENTE_TITULO)

    par_sub = doc.add_paragraph()
    par_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_sub.paragraph_format.space_after = Pt(10)
    _run(par_sub, subtitulo, size=9, italic=True, color=GRIS_TEXTO)

    _tira_info(doc, {'semestre': semestre, 'ciclo': ciclo,
                     'especialidad': especialidad, 'turno': turno})

    # Tabla principal
    ancho_hora = Cm(2.4)
    ancho_dia = Cm((17.0 - 2.4) / len(dias))
    tabla = doc.add_table(rows=1, cols=1 + len(dias))
    tabla.autofit = False

    fila_enc = tabla.rows[0]
    _repetir_encabezado(fila_enc)
    encabezados = ['HORA'] + [str(d).upper() for d in dias]
    for celda, texto in zip(fila_enc.cells, encabezados):
        celda.width = ancho_hora if texto == 'HORA' else ancho_dia
        _fondo_celda(celda, VERDE)
        _borde_celda(celda, VERDE_OSCURO, sz=4)
        _margenes_celda(celda, top=90, bottom=90)
        celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = celda.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(0)
        _run(par, texto, size=9.5, bold=True, color=BLANCO)

    for idx, periodo in enumerate(horarios):
        es_receso = bool(periodo.get('receso'))
        fila = tabla.add_row()
        fila.height = Cm(1.05)
        fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        fondo = FILA_ALT if idx % 2 == 0 else BLANCO
        hora_texto = f"{periodo.get('inicio', '')} – {periodo.get('fin', '')}"

        celda_hora = fila.cells[0]
        celda_hora.width = ancho_hora
        _fondo_celda(celda_hora, ROJO_BG if es_receso else HORA_BG)
        _borde_celda(celda_hora)
        _margenes_celda(celda_hora, top=60, bottom=60)
        celda_hora.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par_h = celda_hora.paragraphs[0]
        par_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par_h.paragraph_format.space_after = Pt(0)
        _run(par_h, hora_texto, size=9.5, bold=True,
             color=ROJO if es_receso else VERDE_OSCURO)

        for celda in fila.cells[1:]:
            celda.width = ancho_dia
            _fondo_celda(celda, ROJO_BG if es_receso else fondo)
            _borde_celda(celda)
            _margenes_celda(celda, top=60, bottom=60)
            celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            par = celda.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            if es_receso:
                _run(par, 'RECESO', size=9, bold=True, color=ROJO)

    # Notas
    par_notas_t = doc.add_paragraph()
    par_notas_t.paragraph_format.space_before = Pt(14)
    par_notas_t.paragraph_format.space_after = Pt(3)
    _run(par_notas_t, 'Notas', size=11, bold=True, color=VERDE,
         font=FUENTE_TITULO)
    _borde_inferior_texto(par_notas_t)

    notas = [
        '1. Este horario es de carácter informativo y puede estar sujeto a cambios.',
        '2. Para cualquier modificación, favor de dirigirse a la Coordinación Académica.',
        f'3. Generado el {T.fecha_larga()} desde la Plataforma CECyTE.',
    ]
    for i, nota in enumerate(notas):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        _run(par, nota, size=8.5 if i < 2 else 8,
             color=GRIS_TEXTO if i < 2 else GRIS_CLARO)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _borde_inferior_texto(par):
    """Filete verde bajo un párrafo (para títulos de sección)."""
    p_pr = par._p.get_or_add_pPr()
    bordes = OxmlElement('w:pBdr')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), '8')
    el.set(qn('w:space'), '2')
    el.set(qn('w:color'), VERDE)
    bordes.append(el)
    p_pr.append(bordes)

