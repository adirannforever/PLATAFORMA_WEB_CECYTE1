# ============================================================
# generador_excel.py — Exportaciones XLSX (openpyxl).
# Encabezado institucional con logo, tablas con estilos de la
# plataforma, formato condicional y configuración de impresión.
# ============================================================
import io
import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

import tema as T

# Paleta (ARGB)
VERDE = 'FF1A6B35'
VERDE_OSCURO = 'FF0F4521'
VERDE_MUTE = 'FFE8F5ED'
FILA_ALT = 'FFF4F9F6'
DORADO = 'FFC8992A'
DORADO_MUTE = 'FFF5E9C8'
GRIS_BORDE = 'FFDCE6E0'
GRIS_TEXTO = 'FF4A6357'
GRIS_CLARO = 'FF94A89D'
GRIS_OSCURO = 'FF1E2E27'
ROJO = 'FFC0392B'
NARANJA = 'FFF37238'
BLANCO = 'FFFFFFFF'

BORDE = Border(*[Side(style='thin', color=GRIS_BORDE)] * 4)


def _fecha_corta(fecha=None):
    f = fecha or date.today()
    return f.strftime('%d/%m/%Y')


def _estilo_encabezado(celda):
    celda.font = Font(name='Calibri', size=11, bold=True, color=BLANCO)
    celda.fill = PatternFill('solid', fgColor=VERDE)
    celda.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
    celda.border = Border(*[Side(style='thin', color=VERDE_OSCURO)] * 4)


def _estilo_dato(celda, *, centrado=False, alterna=False, bold=False,
                 color=GRIS_OSCURO):
    celda.font = Font(name='Calibri', size=11, bold=bold, color=color)
    celda.fill = PatternFill('solid', fgColor=FILA_ALT if alterna else BLANCO)
    celda.alignment = Alignment(horizontal='center' if centrado else 'left',
                                vertical='center')
    celda.border = BORDE


def _encabezado_reporte(ws, titulo, subtitulo, num_columnas):
    """Franja institucional: logo + título + subtítulo (filas 1-3)."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '1A6B35'

    ruta_logo = T.RUTA_LOGO
    if os.path.exists(ruta_logo):
        try:
            img = XLImage(ruta_logo)
            img.width = 40
            img.height = 40
            img.anchor = 'A1'
            ws.add_image(img)
        except Exception:
            pass

    for col in range(1, num_columnas + 1):
        ws.cell(row=1, column=col).fill = PatternFill('solid', fgColor=VERDE_OSCURO)
    ws.merge_cells(start_row=1, start_column=2, end_row=1,
                   end_column=num_columnas)
    celda_titulo = ws.cell(row=1, column=2, value=titulo.upper())
    celda_titulo.font = Font(name='Calibri', size=16, bold=True, color=BLANCO)
    celda_titulo.alignment = Alignment(horizontal='left', vertical='center',
                                       indent=1)
    ws.row_dimensions[1].height = 44

    ws.merge_cells(start_row=2, start_column=1, end_row=2,
                   end_column=num_columnas)
    celda_sub = ws.cell(row=2, column=1, value=subtitulo)
    for col in range(1, num_columnas + 1):
        ws.cell(row=2, column=col).fill = PatternFill('solid', fgColor=VERDE_MUTE)
    celda_sub.font = Font(name='Calibri', size=9, bold=True, color=VERDE)
    celda_sub.alignment = Alignment(horizontal='left', vertical='center',
                                    indent=1)
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 6


def _fila_encabezados(ws, fila, encabezados, anchos):
    for i, (texto, ancho) in enumerate(zip(encabezados, anchos), start=1):
        celda = ws.cell(row=fila, column=i, value=texto)
        _estilo_encabezado(celda)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 22


def _configurar_impresion(ws, ultima_fila, num_columnas, horizontal=False):
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:{get_column_letter(num_columnas)}{ultima_fila}'
    ws.print_title_rows = '4:4'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    margen = 0.5
    ws.page_margins.left = margen
    ws.page_margins.right = margen
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6


# ------------------------------------------------------------
def generar_listado(payload):
    """Listado de alumnos (XLSX)."""
    alumnos = payload.get('alumnos') or []
    inst = payload.get('institucion') or {}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Alumnos'

    _encabezado_reporte(
        ws, 'Listado de alumnos',
        f"{inst.get('nombre', 'CECyTE Plantel 1')} · "
        f"{inst.get('sede', 'Macuspana')} · Exportado el {_fecha_corta()}", 8)

    encabezados = ['Matrícula', 'Apellidos', 'Nombre', 'Semestre', 'Grupo',
                   'Letra', 'Especialidad', 'Estatus']
    anchos = [14, 24, 22, 10, 12, 8, 30, 12]
    _fila_encabezados(ws, 4, encabezados, anchos)

    fila = 5
    for alumno in alumnos:
        estatus = str(alumno.get('estatus', '') or '').lower()
        activo = estatus == 'activo'
        valores = [
            alumno.get('matricula'), alumno.get('apellidos'),
            alumno.get('nombre'), alumno.get('semestre_actual'),
            alumno.get('grupo_nombre'), alumno.get('grupo_letra'),
            alumno.get('especialidad_nombre'), estatus.upper(),
        ]
        alterna = (fila % 2) == 0
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            centrado = col in (1, 4, 5, 6, 8)
            if col == 8:
                _estilo_dato(celda, centrado=True, alterna=alterna, bold=True,
                             color=VERDE if activo else ROJO)
            else:
                _estilo_dato(celda, centrado=centrado, alterna=alterna)
        ws.row_dimensions[fila].height = 18
        fila += 1

    if not alumnos:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        celda = ws.cell(row=5, column=1,
                        value='No hay alumnos que coincidan con los filtros.')
        celda.font = Font(name='Calibri', size=10, italic=True, color=GRIS_CLARO)
        celda.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[5].height = 26
        fila = 6

    _configurar_impresion(ws, fila - 1, 8, horizontal=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------
def generar_estadisticas(payload):
    """Estadísticas por materia (XLSX) con escala de color y barra de datos."""
    stats = payload.get('stats') or []
    inst = payload.get('institucion') or {}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Estadísticas'

    _encabezado_reporte(
        ws, 'Estadísticas académicas',
        f"{inst.get('nombre', 'CECyTE Plantel 1')} · "
        f"Ciclo en curso · Exportado el {_fecha_corta()}", 8)

    encabezados = ['Grupo', 'Letra', 'Materia', 'Total alumnos', 'Promedio',
                   'Aprobados', 'Reprobados', '% Aprobación']
    anchos = [12, 8, 28, 13, 11, 12, 12, 13]
    _fila_encabezados(ws, 4, encabezados, anchos)

    fila = 5
    for fila_dato in stats:
        total = int(fila_dato.get('total_alumnos') or 0)
        aprobados = int(fila_dato.get('aprobados') or 0)
        reprobados = int(fila_dato.get('reprobados') or 0)
        try:
            promedio = round(float(fila_dato.get('promedio_materia') or 0), 2)
        except (TypeError, ValueError):
            promedio = 0
        porcentaje = (aprobados / total) if total > 0 else 0
        alterna = (fila % 2) == 0

        valores = [
            fila_dato.get('grupo_nombre'), fila_dato.get('grupo_letra'),
            fila_dato.get('materia_nombre'), total, promedio,
            aprobados, reprobados, porcentaje,
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            centrado = col != 3
            if col == 7 and reprobados > 0:
                _estilo_dato(celda, centrado=True, alterna=alterna, bold=True,
                             color=ROJO)
            elif col == 5:
                _estilo_dato(celda, centrado=True, alterna=alterna, bold=True,
                             color=VERDE_OSCURO)
            else:
                _estilo_dato(celda, centrado=centrado, alterna=alterna)
        ws.cell(row=fila, column=5).number_format = '0.00'
        ws.cell(row=fila, column=8).number_format = '0%'
        ws.row_dimensions[fila].height = 18
        fila += 1

    ultima = max(fila - 1, 5)
    if fila > 5:
        rango_prom = f'E5:E{fila - 1}'
        ws.conditional_formatting.add(rango_prom, ColorScaleRule(
            start_type='num', start_value=5, start_color=ROJO,
            mid_type='num', mid_value=7, mid_color=DORADO,
            end_type='num', end_value=10, end_color=VERDE))
        rango_pct = f'H5:H{fila - 1}'
        ws.conditional_formatting.add(rango_pct, DataBarRule(
            start_type='num', start_value=0, end_type='num', end_value=1,
            color=VERDE, showValue=True))
    else:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        celda = ws.cell(row=5, column=1,
                        value='No hay estadísticas disponibles para el ciclo.')
        celda.font = Font(name='Calibri', size=10, italic=True, color=GRIS_CLARO)
        celda.alignment = Alignment(horizontal='center', vertical='center')

    _configurar_impresion(ws, ultima, 8, horizontal=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------
def _color_estado(estado):
    e = str(estado or '').lower()
    if 'present' in e or 'asist' in e:
        return VERDE
    if 'retard' in e:
        return NARANJA
    if 'falta' in e or 'inasist' in e:
        return ROJO
    if 'justif' in e:
        return DORADO
    return GRIS_OSCURO


def generar_asistencias(payload):
    """Asistencias de clase (XLSX)."""
    info = payload.get('info') or {}
    asistencias = payload.get('asistencias') or []
    inst = payload.get('institucion') or {}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Asistencias'

    materia = info.get('materia_nombre', '—')
    grupo = f"{info.get('grupo_nombre', '—')} ({info.get('grupo_letra', '')})"
    ciclo = info.get('ciclo_nombre', '—')
    try:
        fecha_clase = datetime.strptime(str(info.get('fecha'))[:10],
                                        '%Y-%m-%d').strftime('%d/%m/%Y')
    except (ValueError, TypeError):
        fecha_clase = str(info.get('fecha') or '—')

    _encabezado_reporte(
        ws, f'Asistencias · {materia}',
        f"Grupo {grupo} · Ciclo {ciclo} · {inst.get('nombre', 'CECyTE Plantel 1')}", 4)

    # Chips de contexto (fila 3): fecha + exportación
    etiquetas = [('FECHA DE CLASE', fecha_clase),
                 ('EXPORTADO', datetime.now().strftime('%d/%m/%Y %H:%M'))]
    for i, (etiqueta, valor) in enumerate(etiquetas):
        col = 1 + i * 2
        celda_etq = ws.cell(row=3, column=col, value=etiqueta)
        celda_etq.font = Font(name='Calibri', size=8, bold=True, color=VERDE)
        celda_etq.fill = PatternFill('solid', fgColor=VERDE_MUTE)
        celda_etq.alignment = Alignment(horizontal='right', vertical='center')
        celda_val = ws.cell(row=3, column=col + 1, value=valor)
        celda_val.font = Font(name='Calibri', size=9, color=GRIS_OSCURO)
        celda_val.fill = PatternFill('solid', fgColor=VERDE_MUTE)
        celda_val.alignment = Alignment(horizontal='left', vertical='center',
                                        indent=1)
    ws.row_dimensions[3].height = 16

    encabezados = ['Alumno', 'Matrícula', 'Estado', 'Justificación']
    anchos = [34, 15, 15, 28]
    _fila_encabezados(ws, 4, encabezados, anchos)

    fila = 5
    for alumno in asistencias:
        nombre_completo = (f"{alumno.get('apellidos', '')}, "
                           f"{alumno.get('nombre', '')}")
        estado = alumno.get('estado', '')
        alterna = (fila % 2) == 0
        valores = [nombre_completo, alumno.get('matricula') or 'N/A',
                   estado, alumno.get('justificacion') or '']
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=col, value=valor)
            if col == 3:
                _estilo_dato(celda, centrado=True, alterna=alterna, bold=True,
                             color=_color_estado(estado))
            else:
                _estilo_dato(celda, centrado=(col == 2), alterna=alterna)
        ws.row_dimensions[fila].height = 18
        fila += 1

    if not asistencias:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=4)
        celda = ws.cell(row=5, column=1,
                        value='No hay asistencias registradas para esta fecha.')
        celda.font = Font(name='Calibri', size=10, italic=True, color=GRIS_CLARO)
        celda.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[5].height = 26
        fila = 6

    _configurar_impresion(ws, fila - 1, 4, horizontal=False)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()



